# app.py
import streamlit as st
from pbixray.core import PBIXRay
import os
import tempfile
from docx import Document
from docx.shared import Inches
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from io import BytesIO
import pandas as pd

def generate_word_doc(report_data):
    """Generates a Word document from the extracted report data."""
    document = Document()

    document.add_heading('Power BI Report Documentation', 0)

    # Add Metadata
    document.add_heading('Metadata', level=1)
    if report_data.get("metadata"):
        for key, value in report_data["metadata"].items():
            document.add_paragraph(f"{key}: {value}")
    else:
        document.add_paragraph("No metadata available.")

    # Add Schema
    document.add_heading('Schema', level=1)
    if report_data.get("schema"):
        for table in report_data["schema"]:
            document.add_heading(f"Table: {table.get('name', 'N/A')}", level=2)
            if table.get("columns"):
                document.add_paragraph("Columns:")
                for column in table["columns"]:
                    document.add_paragraph(f"- {column.get('name', 'N/A')} ({column.get('dataType', 'N/A')})")
            else:
                document.add_paragraph("No columns found for this table.")
    else:
        document.add_paragraph("No schema information available.")

    # Add Relationships
    document.add_heading('Relationships', level=1)
    if report_data.get("relationships") is not None and not report_data.get("relationships").empty:
        for rel in report_data["relationships"].to_dict('records'):
            document.add_paragraph(f"From Table: {rel.get('fromTable', 'N/A')}, From Column: {rel.get('fromColumn', 'N/A')}, To Table: {rel.get('toTable', 'N/A')}, To Column: {rel.get('toColumn', 'N/A')}")
    else:
        document.add_paragraph("No relationships available.")

    # Add Power Query Code
    document.add_heading('Power Query Code', level=1)
    if report_data.get("power_query") is not None and not report_data.get("power_query").empty:
        for pq in report_data["power_query"].to_dict('records'):
             document.add_paragraph(f"Name: {pq.get('name', 'N/A')}")
             if pq.get('expression'):
                  document.add_paragraph("Expression:")
                  document.add_paragraph(pq['expression'])
    else:
        document.add_paragraph("No Power Query code available.")

    # Add M Parameters
    document.add_heading('M Parameters', level=1)
    if report_data.get("m_parameters") is not None and not report_data.get("m_parameters").empty:
        for param in report_data["m_parameters"].to_dict('records'):
            document.add_paragraph(f"Name: {param.get('name', 'N/A')}, Value: {param.get('value', 'N/A')}")
    else:
        document.add_paragraph("No M parameters available.")

    # Add DAX Tables
    document.add_heading('DAX Tables', level=1)
    if report_data.get("dax_tables") is not None and not report_data.get("dax_tables").empty:
         for table in report_data["dax_tables"].to_dict('records'):
            document.add_paragraph(f"Name: {table.get('name', 'N/A')}")
            if table.get('expression'):
                 document.add_paragraph("Expression:")
                 document.add_paragraph(table['expression'])
    else:
        document.add_paragraph("No DAX tables available.")

    # Add DAX Measures
    document.add_heading('DAX Measures', level=1)
    if report_data.get("dax_measures") is not None and not report_data.get("dax_measures").empty:
        for measure in report_data["dax_measures"].to_dict('records'):
            document.add_paragraph(f"Name: {measure.get('name', 'N/A')}")
            if measure.get('expression'):
                document.add_paragraph("Expression:")
                document.add_paragraph(measure['expression'])
    else:
        document.add_paragraph("No DAX measures available.")

    document_stream = BytesIO()
    document.save(document_stream)
    document_stream.seek(0)
    return document_stream

def generate_pdf_doc(report_data):
    """Generates a PDF document from the extracted report data."""
    buffer = BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    width, height = letter

    c.drawString(100, height - 50, "Power BI Report Documentation")
    y_position = height - 100

    def draw_text(text, x, y, size=12, bold=False):
        c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
        c.drawString(x, y, text)
        return y - size - 5 # Return new y position

    def draw_paragraph(text, x, y, size=10):
        c.setFont("Helvetica", size)
        # Simple text wrapping (can be improved)
        from reportlab.platypus import SimpleDocTemplate, Paragraph
        from reportlab.lib.styles import getSampleStyleSheet
        styles = getSampleStyleSheet()
        style = styles['Normal']
        story = [Paragraph(text, style)]
        # This simple approach doesn't work well directly on canvas,
        # but for basic text we can just draw line by line if needed
        # For this example, we'll just draw a single line for simplicity.
        # A more robust solution would use ReportLab's flowables.
        c.drawString(x, y, text[:100] + "..." if len(text) > 100 else text) # Basic truncation for simplicity
        return y - size - 5


    # Add Metadata
    y_position = draw_text("Metadata", 100, y_position, size=14, bold=True)
    if report_data.get("metadata"):
        for key, value in report_data["metadata"].items():
            y_position = draw_paragraph(f"{key}: {value}", 100, y_position)
    else:
        y_position = draw_paragraph("No metadata available.", 100, y_position)

    # Add Schema
    y_position = draw_text("Schema", 100, y_position, size=14, bold=True)
    if report_data.get("schema"):
        for table in report_data["schema"]:
            y_position = draw_text(f"Table: {table.get('name', 'N/A')}", 100, y_position, size=12, bold=True)
            if table.get("columns"):
                y_position = draw_text("Columns:", 100, y_position, size=10)
                for column in table["columns"]:
                     y_position = draw_paragraph(f"- {column.get('name', 'N/A')} ({column.get('dataType', 'N/A')})", 100, y_position)
            else:
                 y_position = draw_paragraph("No columns found for this table.", 100, y_position)
    else:
        y_position = draw_paragraph("No schema information available.", 100, y_position)

    # Add Relationships
    y_position = draw_text("Relationships", 100, y_position, size=14, bold=True)
    if report_data.get("relationships") is not None and not report_data.get("relationships").empty:
        for rel in report_data["relationships"].to_dict('records'):
            y_position = draw_paragraph(f"From Table: {rel.get('fromTable', 'N/A')}, From Column: {rel.get('fromColumn', 'N/A')}, To Table: {rel.get('toTable', 'N/A')}, To Column: {rel.get('toColumn', 'N/A')}", 100, y_position)
    else:
        y_position = draw_paragraph("No relationships available.", 100, y_position)

    # Add Power Query Code
    y_position = draw_text("Power Query Code", 100, y_position, size=14, bold=True)
    if report_data.get("power_query") is not None and not report_data.get("power_query").empty:
        for pq in report_data["power_query"].to_dict('records'):
            y_position = draw_paragraph(f"Name: {pq.get('name', 'N/A')}", 100, y_position)
            if pq.get('expression'):
                y_position = draw_text("Expression:", 100, y_position, size=10)
                y_position = draw_paragraph(pq['expression'], 100, y_position)
    else:
        y_position = draw_paragraph("No Power Query code available.", 100, y_position)

    # Add M Parameters
    y_position = draw_text("M Parameters", 100, y_position, size=14, bold=True)
    if report_data.get("m_parameters") is not None and not report_data.get("m_parameters").empty:
        for param in report_data["m_parameters"].to_dict('records'):
            y_position = draw_paragraph(f"Name: {param.get('name', 'N/A')}, Value: {param.get('value', 'N/A')}", 100, y_position)
    else:
        y_position = draw_paragraph("No M parameters available.", 100, y_position)

    # Add DAX Tables
    y_position = draw_text("DAX Tables", 100, y_position, size=14, bold=True)
    if report_data.get("dax_tables") is not None and not report_data.get("dax_tables").empty:
         for table in report_data["dax_tables"].to_dict('records'):
            y_position = draw_paragraph(f"Name: {table.get('name', 'N/A')}", 100, y_position)
            if table.get('expression'):
                 y_position = draw_text("Expression:", 100, y_position, size=10)
                 y_position = draw_paragraph(table['expression'], 100, y_position)
    else:
        y_position = draw_paragraph("No DAX tables available.", 100, y_position)

    # Add DAX Measures
    y_position = draw_text("DAX Measures", 100, y_position, size=14, bold=True)
    if report_data.get("dax_measures") is not None and not report_data.get("dax_measures").empty:
        for measure in report_data["dax_measures"].to_dict('records'):
            y_position = draw_paragraph(f"Name: {measure.get('name', 'N/A')}", 100, y_position)
            if measure.get('expression'):
                y_position = draw_text("Expression:", 100, y_position, size=10)
                y_position = draw_paragraph(measure['expression'], 100, y_position)
    else:
        y_position = draw_paragraph("No DAX measures available.", 100, y_position)

    c.save()
    buffer.seek(0)
    return buffer


def main():
    st.title("Power BI Report Documentation Generator")

    uploaded_file = st.file_uploader("Upload your Power BI .pbix file", type="pbix")

    if uploaded_file is not None:
        try:
            # Create a temporary file to save the uploaded pbix
            with tempfile.NamedTemporaryFile(delete=False, suffix=".pbix") as tmp_file:
                tmp_file.write(uploaded_file.getvalue())
                tmp_pbix_path = tmp_file.name

            st.success(f"File uploaded successfully: {uploaded_file.name}")

            # Initialize PBIXRay with the temporary file path
            pbix_ray = PBIXRay(tmp_pbix_path)

            st.subheader("Extracting Report Information:")

            # Extract various pieces of information using pbixray
            metadata = pbix_ray.metadata
            st.write("Metadata:", metadata)

            schema = pbix_ray.schema
            st.write("Schema:", schema)

            relationships = pbix_ray.relationships
            st.write("Relationships:", relationships)

            power_query = pbix_ray.power_query
            st.write("Power Query:", power_query)

            m_parameters = pbix_ray.m_parameters
            st.write("M Parameters:", m_parameters)

            dax_tables = pbix_ray.dax_tables
            st.write("DAX Tables:", dax_tables)

            dax_measures = pbix_ray.dax_measures
            st.write("DAX Measures:", dax_measures)


            st.success("Information extracted successfully!")

            # Store extracted information in a dictionary for later use
            report_data = {
                "metadata": metadata,
                "schema": schema,
                "relationships": relationships,
                "power_query": power_query,
                "m_parameters": m_parameters,
                "dax_tables": dax_tables,
                "dax_measures": dax_measures,
            }

            # Clean up the temporary file
            os.remove(tmp_pbix_path)

            st.subheader("Download Documentation:")

            # Add download buttons for Word and PDF
            word_doc_stream = generate_word_doc(report_data)
            st.download_button(
                label="Download as Word (.docx)",
                data=word_doc_stream,
                file_name=f"{os.path.splitext(uploaded_file.name)[0]}_documentation.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

            pdf_doc_stream = generate_pdf_doc(report_data)
            st.download_button(
                label="Download as PDF (.pdf)",
                data=pdf_doc_stream,
                file_name=f"{os.path.splitext(uploaded_file.name)[0]}_documentation.pdf",
                mime="application/pdf"
            )


        except Exception as e:
            st.error(f"An error occurred: {e}")
            # Ensure temporary file is removed even if an error occurs
            if 'tmp_pbix_path' in locals() and os.path.exists(tmp_pbix_path):
                os.remove(tmp_pbix_path)


if __name__ == "__main__":
    main()
